"""Outline a PowerPoint deck and flag layout and accessibility problems.

    python -B pptx_inspect.py <file> [--json]

Per slide: layout, title, each shape's kind and text, and the speaker notes. Flags: a missing
title, explicit text under 12 pt, shapes outside the slide, pictures without alt text, empty
placeholders, and text likely to overflow its box.
Exit codes: 0 clean, 10 flags raised, 1 error.
"""
from __future__ import annotations

import json
import math
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Emu

TITLE_TYPES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE)

MIN_PT = 12
DEFAULT_PT = 18
EMU_PER_INCH = 914400


def _iter_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == 6:  # group
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _kind(shape) -> str:
    if shape.is_placeholder:
        return f"placeholder:{shape.placeholder_format.type.name.lower()}"
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return "chart"
    if getattr(shape, "has_table", False) and shape.has_table:
        return "table"
    if shape.shape_type == 13:
        return "picture"
    return "text" if shape.has_text_frame else "shape"


def _estimated_overflow(shape) -> bool:
    if not shape.has_text_frame or not shape.width or not shape.height:
        return False
    lines, tallest = 0, DEFAULT_PT
    width_pt = shape.width / EMU_PER_INCH * 72
    for p in shape.text_frame.paragraphs:
        sizes = [r.font.size.pt for r in p.runs if r.font.size] or [DEFAULT_PT]
        size = max(sizes)
        tallest = max(tallest, size)
        chars_per_line = max(1, int(width_pt / (size * 0.5)))
        lines += max(1, math.ceil(len(p.text) / chars_per_line)) * size * 1.2
    return lines > shape.height / EMU_PER_INCH * 72 * 1.05 and len(shape.text_frame.text) > 0


def outline(path: Path) -> dict:
    prs = Presentation(str(path))
    sw, sh = prs.slide_width, prs.slide_height
    slides, flags = [], []
    for i, slide in enumerate(prs.slides, 1):
        # By placeholder type: slide.shapes.title expects index 0, and PptxGenJS numbers from 100.
        title_shape = next((s for s in slide.placeholders if s.placeholder_format.type in TITLE_TYPES), None)
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        if not title:
            flags.append(f"slide {i}: no title (screen readers and the outline use it)")
        shapes = []
        for shape in _iter_shapes(slide.shapes):
            kind = _kind(shape)
            text = shape.text_frame.text.strip() if shape.has_text_frame else ""
            entry = {"kind": kind, "name": shape.name, "text": text[:200]}
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size is not None and r.font.size.pt < MIN_PT and r.text.strip():
                            flags.append(f"slide {i}: '{shape.name}' has {r.font.size.pt:g} pt text")
                            break
                if shape.is_placeholder and not text:
                    flags.append(f"slide {i}: empty placeholder '{shape.name}'")
                if _estimated_overflow(shape):
                    flags.append(f"slide {i}: text in '{shape.name}' likely overflows its box")
            if shape.left is not None and shape.width is not None:
                if shape.left < 0 or shape.top < 0 or shape.left + shape.width > sw or shape.top + shape.height > sh:
                    flags.append(f"slide {i}: '{shape.name}' extends outside the slide")
            if kind in ("picture", "chart"):
                frame = "./p:nvPicPr/p:cNvPr" if kind == "picture" else "./p:nvGraphicFramePr/p:cNvPr"
                descr = shape._element.xpath(frame)[0].get("descr") or ""
                entry["alt"] = descr
                if not descr.strip():
                    flags.append(f"slide {i}: {kind} '{shape.name}' has no alt text")
            shapes.append(entry)
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        slides.append({"slide": i, "layout": slide.slide_layout.name, "title": title, "shapes": shapes, "notes": notes})
    return {"file": str(path), "widthIn": round(Emu(sw).inches, 3), "heightIn": round(Emu(sh).inches, 3),
            "slides": slides, "flags": flags}


def main(argv: list[str]) -> int:
    if not argv:
        print(__doc__.strip(), file=sys.stderr)
        return 1
    try:
        result = outline(Path(argv[0]))
    except (OSError, ValueError, KeyError) as e:
        print(f"error: {e}", file=sys.stderr)
        return 1
    if "--json" in argv:
        print(json.dumps(result, indent=2))
    else:
        print(f"{len(result['slides'])} slides, {result['widthIn']} x {result['heightIn']} in")
        for s in result["slides"]:
            print(f"[{s['slide']}] {s['title'] or '(no title)'}  ({s['layout']})")
            for sh in s["shapes"]:
                if sh["kind"] in ("placeholder:title", "placeholder:center_title", "placeholder:vertical_title"):
                    continue
                print(f"    {sh['kind']}: {sh['text'][:80] or sh.get('alt', '')}")
            if s["notes"]:
                print(f"    notes: {s['notes'][:80]}")
        for f in result["flags"]:
            print(f"FLAG {f}")
    return 10 if result["flags"] else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
