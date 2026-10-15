"""Inspect an Excel workbook without Excel.

    python -B xlsx_inspect.py summary <file> [--sheet NAME] [--max-rows N] [--json]
    python -B xlsx_inspect.py check   <file>

summary lists sheets, dimensions, headers, sample rows (cached values), formulas, charts, and names.
check reports error values and formulas that point at sheets that do not exist.
Exit codes: 0 clean, 10 findings, 1 error.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer

ERROR_VALUES = {"#REF!", "#DIV/0!", "#NAME?", "#N/A", "#VALUE!", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}


def _cell_text(v) -> str | None:
    if v is None:
        return None
    return v.isoformat() if hasattr(v, "isoformat") else v if isinstance(v, (int, float, bool)) else str(v)


def summary(path: Path, sheet: str | None = None, max_rows: int = 5) -> dict:
    formulas = load_workbook(path)
    values = load_workbook(path, data_only=True)
    out = {"file": str(path), "sheets": [], "definedNames": sorted(formulas.defined_names.keys())}
    for ws in formulas.worksheets:
        if sheet and ws.title != sheet:
            continue
        wv = values[ws.title]
        formula_cells = [c.coordinate for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
        uncached = [ref for ref in formula_cells if wv[ref].value is None]
        rows = list(wv.iter_rows(min_row=1, max_row=min(ws.max_row, max_rows + 1), values_only=True))
        out["sheets"].append({
            "title": ws.title,
            "state": ws.sheet_state,
            "dimensions": ws.dimensions,
            "rows": ws.max_row,
            "columns": ws.max_column,
            "freezePanes": ws.freeze_panes,
            "autoFilter": ws.auto_filter.ref,
            "header": [_cell_text(v) for v in rows[0]] if rows else [],
            "sample": [[_cell_text(v) for v in r] for r in rows[1:]],
            "formulas": len(formula_cells),
            "formulasWithoutCachedValue": len(uncached),
            "mergedRanges": [str(r) for r in ws.merged_cells.ranges],
            "charts": [type(c).__name__ for c in ws._charts],
            "images": len(ws._images),
            "dataValidations": len(ws.data_validations.dataValidation),
            "conditionalFormats": len(ws.conditional_formatting),
        })
    return out


def check(path: Path) -> list[str]:
    formulas = load_workbook(path)
    values = load_workbook(path, data_only=True)
    sheets = set(formulas.sheetnames)
    findings = []
    for ws in formulas.worksheets:
        wv = values[ws.title]
        for row in ws.iter_rows():
            for c in row:
                cached = wv[c.coordinate].value
                if isinstance(cached, str) and cached in ERROR_VALUES:
                    findings.append(f"{ws.title}!{c.coordinate}: {cached}")
                if not (isinstance(c.value, str) and c.value.startswith("=")):
                    continue
                if "#REF!" in c.value:
                    findings.append(f"{ws.title}!{c.coordinate}: formula contains #REF! ({c.value})")
                    continue
                try:
                    tokens = Tokenizer(c.value).items
                except Exception as e:  # the tokenizer raises bare exceptions on malformed input
                    findings.append(f"{ws.title}!{c.coordinate}: formula does not parse ({e})")
                    continue
                for t in tokens:
                    if t.type == "OPERAND" and t.subtype == "RANGE" and "!" in t.value:
                        target = t.value.rsplit("!", 1)[0]
                        target = re.sub(r"^\[\d+\]", "", target.strip("'").replace("''", "'"))
                        if target not in sheets and not t.value.startswith("["):
                            findings.append(f"{ws.title}!{c.coordinate}: refers to missing sheet '{target}' ({c.value})")
    return findings


def main(argv: list[str]) -> int:
    if len(argv) < 2 or argv[0] not in ("summary", "check"):
        print(__doc__.strip(), file=sys.stderr)
        return 1
    path = Path(argv[1])
    try:
        if argv[0] == "summary":
            sheet = argv[argv.index("--sheet") + 1] if "--sheet" in argv else None
            max_rows = int(argv[argv.index("--max-rows") + 1]) if "--max-rows" in argv else 5
            result = summary(path, sheet, max_rows)
            if "--json" in argv:
                print(json.dumps(result, indent=2, default=str))
            else:
                for s in result["sheets"]:
                    print(f"[{s['title']}] {s['dimensions']} ({s['rows']} rows x {s['columns']} cols), "
                          f"{s['formulas']} formulas ({s['formulasWithoutCachedValue']} without a cached value), "
                          f"charts: {', '.join(s['charts']) or 'none'}")
                    print(f"  header: {s['header']}")
                    for r in s["sample"]:
                        print(f"  {r}")
                if result["definedNames"]:
                    print(f"defined names: {', '.join(result['definedNames'])}")
            return 0
        findings = check(path)
        for f in findings:
            print(f)
        print(f"{len(findings)} finding(s)" if findings else "No error values or broken sheet references.")
        return 10 if findings else 0
    except (OSError, KeyError, ValueError) as e:
        print(f"error: {e}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
