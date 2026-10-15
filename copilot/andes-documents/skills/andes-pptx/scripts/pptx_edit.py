"""Edit an existing PowerPoint deck: replace text, set speaker notes, set picture or chart alt text.

    python -B pptx_edit.py replace  <in> <out> --map <mapping.json>
    python -B pptx_edit.py notes    <in> <out> --slide N --text <notes>
    python -B pptx_edit.py alt-text <in> <out> --slide N --shape <name> --text <description>

replace swaps every key of a JSON object for its value in every text frame, table cell, and
group on every slide, and in the speaker notes. Text split across runs is still found; the
replacement takes the formatting of the run where the match starts.
Exit codes: 0 done, 10 nothing matched, 1 error.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Iterator

from pptx import Presentation


def _iter_shapes(shapes) -> Iterator:
    """Every shape, descending into groups, in the same order pptx_inspect reports them."""
    for shape in shapes:
        if shape.shape_type == 6:  # group
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _text_frames(shapes) -> Iterator:
    for shape in shapes:
        if shape.shape_type == 6:  # group
            yield from _text_frames(shape.shapes)
        elif getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame
        elif shape.has_text_frame:
            yield shape.text_frame


def replace_in_runs(runs, key: str, value: str) -> int:
    if not key:
        raise ValueError("an empty key matches everywhere; every mapping key must be non-empty")
    count, offset = 0, 0
    while True:
        texts = [r.text for r in runs]
        full = "".join(texts)
        i = full.find(key, offset)
        if i < 0:
            return count
        end, pos, first = i + len(key), 0, True
        for r, t in zip(runs, texts):
            start, stop = pos, pos + len(t)
            pos = stop
            if stop <= i or start >= end:
                continue
            a, b = max(i, start) - start, min(end, stop) - start
            r.text = t[:a] + (value if first else "") + t[b:]
            first = False
        count += 1
        offset = i + len(value)


def replace_text(prs, mapping: dict[str, str]) -> dict[str, int]:
    if not isinstance(mapping, dict) or any(not isinstance(k, str) or not k for k in mapping):
        raise ValueError("the mapping must be a JSON object whose keys are non-empty strings")
    counts = {k: 0 for k in mapping}
    for slide in prs.slides:
        frames = list(_text_frames(slide.shapes))
        if slide.has_notes_slide:
            frames.append(slide.notes_slide.notes_text_frame)
        for frame in frames:
            for p in frame.paragraphs:
                for key, value in mapping.items():
                    if key in "".join(r.text for r in p.runs):
                        counts[key] += replace_in_runs(p.runs, key, str(value))
    return counts


def main(argv: list[str]) -> int:
    if len(argv) < 3 or argv[0] not in ("replace", "notes", "alt-text"):
        print(__doc__.strip(), file=sys.stderr)
        return 1
    command, src, dst = argv[0], Path(argv[1]), Path(argv[2])
    if src.resolve() == dst.resolve():
        print("error: write to a new file; the input is never overwritten", file=sys.stderr)
        return 1
    opt = lambda name, default=None: argv[argv.index(name) + 1] if name in argv else default  # noqa: E731
    try:
        prs = Presentation(str(src))
        if command == "replace":
            with open(opt("--map"), encoding="utf-8") as f:
                counts = replace_text(prs, json.load(f))
            for k, n in counts.items():
                print(f"{n:>4}  {k}")
            matched = sum(counts.values())
        else:
            index, text = int(opt("--slide", "0")), opt("--text")
            if not 1 <= index <= len(prs.slides) or text is None:
                print(f"error: --slide must be 1-{len(prs.slides)} and --text is required", file=sys.stderr)
                return 1
            slide = prs.slides[index - 1]
            if command == "notes":
                slide.notes_slide.notes_text_frame.text = text
                matched = 1
            else:
                name = opt("--shape")
                targets = []
                for s in _iter_shapes(slide.shapes):
                    if name is not None and s.name != name:
                        continue
                    if s.shape_type == 13:
                        targets.append(s._element.xpath("./p:nvPicPr/p:cNvPr")[0])
                    elif getattr(s, "has_chart", False) and s.has_chart:
                        targets.append(s._element.xpath("./p:nvGraphicFramePr/p:cNvPr")[0])
                for c_nv_pr in targets:
                    c_nv_pr.set("descr", text)
                matched = len(targets)
                print(f"{matched} picture(s) or chart(s) updated")
        prs.save(str(dst))
        return 0 if matched else 10
    except (OSError, ValueError, KeyError, TypeError) as e:
        print(f"error: {e}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
